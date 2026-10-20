"""使用 standard_business_brief 预设生成 DOCX 质检报告。"""

from pathlib import Path
from typing import Any, NamedTuple

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from smart_compliance.domain.models import Evidence, ReviewItem, ReviewReport, Verdict
from smart_compliance.reporting.audit import public_raw_output_file, source_hash

FONT_NAME = "Microsoft YaHei"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
RISK_RED = RGBColor(155, 28, 28)
CAUTION = RGBColor(122, 90, 0)
INK = RGBColor(34, 34, 34)
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


class StyleToken(NamedTuple):
    """standard_business_brief 段落样式 token。"""

    size: float
    color: RGBColor
    before: float
    after: float
    line: float


def export_docx(report: ReviewReport, output_path: Path) -> Path:
    """导出可复核的 DOCX 报告。

    Args:
        report: 质检报告。
        output_path: 输出路径。

    Returns:
        已生成文件路径。
    """
    document = Document()
    _configure_document(document)
    _add_masthead(document, report)
    _add_summary(document, report)
    _add_items(document, report.items)
    _add_parsing_quality(document, report)
    _add_audit(document, report)
    _add_limitations(document, report)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def _configure_document(document: DocumentObject) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = document.styles
    _configure_style(styles["Normal"], StyleToken(10.2, INK, 0, 5, 1.05))
    _configure_style(styles["Title"], StyleToken(23, INK, 0, 4, 1.0))
    _configure_style(styles["Subtitle"], StyleToken(12, MUTED, 0, 16, 1.0))
    _configure_style(styles["Heading 1"], StyleToken(16, BLUE, 16, 8, 1.0))
    _configure_style(styles["Heading 2"], StyleToken(13, BLUE, 12, 6, 1.0))
    _configure_style(styles["Heading 3"], StyleToken(12, DARK_BLUE, 8, 4, 1.0))
    _set_running_furniture(section)


def _configure_style(style: Any, token: StyleToken) -> None:
    font = style.font
    font.name = FONT_NAME
    font.size = Pt(token.size)
    font.color.rgb = token.color
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(token.before)
    paragraph.space_after = Pt(token.after)
    paragraph.line_spacing = token.line


def _set_running_furniture(section: Any) -> None:
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("智审合规助手  |  采购合同质检报告")
    _set_run(run, 9, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("教学模拟/测试数据  |  AI 结果须人工复核")
    _set_run(run, 8.5, MUTED)


def _add_masthead(document: DocumentObject, report: ReviewReport) -> None:
    title = document.add_paragraph(style="Title")
    title.add_run("采购合同智能质检报告")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(report.project_name)
    metadata = [
        ("报告编号", report.report_id),
        ("样本编号", report.sample_id or "未识别"),
        ("生成时间", report.generated_at.astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")),
        ("质检引擎", f"{report.model.provider} / {report.model.model_name}"),
        ("提示词版本", report.model.prompt_version),
        ("运行 ID", report.model.run_id),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}：")
        _set_run(label_run, 10.5, INK, bold=True)
        _set_run(paragraph.add_run(value), 10.5, INK)
    notice = document.add_paragraph()
    notice.paragraph_format.space_before = Pt(10)
    notice.paragraph_format.space_after = Pt(12)
    _shade_paragraph(notice, "FFF4E5")
    _set_run(notice.add_run("重要提示："), 10.5, CAUTION, bold=True)
    _set_run(notice.add_run(report.disclaimer), 10.5, INK)


def _add_summary(document: DocumentObject, report: ReviewReport) -> None:
    document.add_heading("一、质检摘要", level=1)
    table = document.add_table(rows=2, cols=4)
    second_label = "覆盖编码" if report.summary.covered_clause_count is not None else "一致项"
    second_value = (
        report.summary.covered_clause_count
        if report.summary.covered_clause_count is not None
        else report.summary.consistent_count
    )
    labels = ("总体风险", second_label, "缺失项", "差异项")
    values = (
        report.summary.overall_risk.value.upper(),
        str(second_value),
        str(report.summary.missing_count),
        str(report.summary.different_count),
    )
    for column, label in enumerate(labels):
        _set_cell_text(table.cell(0, column), label, bold=True, color=INK, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(table.cell(0, column), "F2F4F7")
        value_color = RISK_RED if column in {2, 3} and values[column] != "0" else DARK_BLUE
        _set_cell_text(
            table.cell(1, column), values[column], bold=True, color=value_color, alignment=WD_ALIGN_PARAGRAPH.CENTER
        )
    _set_table_geometry(table, [2340, 2340, 2340, 2340])
    document.add_paragraph(
        "结论仅反映系统在当前解析文本和规则/模型配置下识别到的条款承接情况，所有异常项均需回到原始 PDF 复核。"
    )


def _add_items(document: DocumentObject, items: list[ReviewItem]) -> None:
    document.add_heading("二、逐条质检结果", level=1)
    for item in items:
        heading = document.add_heading(f"{item.source_order}. {item.clause_name}", level=2)
        status = heading.add_run(f"  [{_verdict_label(item.verdict)}]")
        _set_run(status, 11, _verdict_color(item.verdict), bold=True)
        _add_labeled_paragraph(document, "采购结果证据", _evidence_text(item.procurement_evidence))
        _add_labeled_paragraph(document, "合同证据", _evidence_text(item.contract_evidence) or "未找到对应证据")
        _add_labeled_paragraph(document, "判断", item.difference)
        _add_labeled_paragraph(document, "风险说明", item.risk_reason)
        _add_labeled_paragraph(document, "复核建议", item.recommendation)
        _add_labeled_paragraph(
            document,
            "可信度/复核",
            f"{item.confidence:.0%}；{'需要人工复核' if item.needs_human_review else '建议抽样复核'}",
        )


def _add_parsing_quality(document: DocumentObject, report: ReviewReport) -> None:
    document.add_heading("三、PDF 解析质量", level=1)
    table = document.add_table(rows=1, cols=6)
    headers = ("文档", "页数", "原始字符", "清洗字符", "移除水印", "低质量页")
    for index, header in enumerate(headers):
        _set_cell_text(table.cell(0, index), header, bold=True, color=INK, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(table.cell(0, index), "F2F4F7")
    for role, quality in report.parsing_quality.items():
        cells = table.add_row().cells
        values = (
            "采购结果" if role == "procurement_result" else "采购合同",
            str(quality.page_count),
            str(quality.raw_characters),
            str(quality.cleaned_characters),
            str(quality.removed_watermark_candidates),
            "、".join(str(page) for page in quality.low_quality_pages) or "无",
        )
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _set_table_geometry(table, [1800, 1000, 1500, 1500, 1500, 2060])


def _add_limitations(document: DocumentObject, report: ReviewReport) -> None:
    document.add_heading("五、局限与人工复核", level=1)
    for index, limitation in enumerate(report.limitations, start=1):
        paragraph = _add_labeled_paragraph(document, f"局限 {index}", limitation)
        paragraph.paragraph_format.space_after = Pt(2)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    _set_run(paragraph.add_run("最终结论："), 10.2, INK, bold=True)
    _set_run(
        paragraph.add_run("请由采购、业务与法务人员结合原始文件完成复核和签章。"),
        10.2,
        RISK_RED,
        bold=True,
    )


def _add_audit(document: DocumentObject, report: ReviewReport) -> None:
    document.add_heading("四、运行与证据审计", level=1)
    result_type = "外部模型辅助结果" if report.model.raw_output_file else "离线规则基线结果"
    raw_output_file = public_raw_output_file(report)
    rows = [
        ("结果类型", result_type),
        ("证据校验", report.evidence_validation_status.value),
        ("人工复核", "需要" if report.summary.needs_human_review else "建议抽样复核"),
        ("采购结果 PDF SHA-256", source_hash(report, "procurement_result_pdf")),
        ("采购合同 PDF SHA-256", source_hash(report, "contract_pdf")),
        ("采购结果 OCR 文本 SHA-256", source_hash(report, "procurement_result_ocr_text")),
        ("采购合同 OCR 文本 SHA-256", source_hash(report, "contract_ocr_text")),
        ("原始模型输出", raw_output_file or "不适用（离线规则未调用 LLM）"),
        ("API 路径", report.model.api_path or "不适用"),
        ("重试次数", str(report.model.retry_count)),
    ]
    for label, value in rows:
        _add_labeled_paragraph(document, label, value)
    if report.validation_warnings:
        for index, warning in enumerate(report.validation_warnings, start=1):
            _add_labeled_paragraph(document, f"校验警告 {index}", warning)
    else:
        _add_labeled_paragraph(document, "校验摘要", "未记录证据校验警告。")


def _add_labeled_paragraph(document: DocumentObject, label: str, text: str) -> Paragraph:
    paragraph = document.add_paragraph()
    _set_run(paragraph.add_run(f"{label}："), 10, INK, bold=True)
    _set_run(paragraph.add_run(text), 10, INK)
    return paragraph


def _evidence_text(evidence: Evidence) -> str:
    if not evidence.quote:
        return ""
    locator = f"第 {evidence.page} 页 / {evidence.chunk_id}" if evidence.page else "页码未知"
    return f"{evidence.quote}（{locator}）"


def _set_run(run: Run, size: float, color: RGBColor, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run_properties = run._element.get_or_add_rPr()
    run_properties.get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)


def _set_cell_text(
    cell: _Cell,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor = INK,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _set_run(paragraph.add_run(text), 9.5, color, bold=bold)


def _set_table_geometry(table: Table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        message = f"表格列宽总和必须为 {TABLE_WIDTH_DXA}: {widths}"
        raise ValueError(message)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    _set_xml_width(properties, "tblW", TABLE_WIDTH_DXA)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            _set_xml_width(cell._tc.get_or_add_tcPr(), "tcW", width)
            _set_cell_margins(cell)


def _set_xml_width(parent: Any, tag: str, width: int) -> None:
    element = parent.find(qn(f"w:{tag}"))
    if element is None:
        element = OxmlElement(f"w:{tag}")
        parent.append(element)
    element.set(qn("w:w"), str(width))
    element.set(qn("w:type"), "dxa")


def _set_cell_margins(cell: _Cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _set_cell_shading(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _shade_paragraph(paragraph: Paragraph, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def _verdict_label(verdict: Verdict) -> str:
    return {Verdict.CONSISTENT: "一致项", Verdict.MISSING: "缺失项", Verdict.DIFFERENT: "差异项"}[verdict]


def _verdict_color(verdict: Verdict) -> RGBColor:
    return DARK_BLUE if verdict == Verdict.CONSISTENT else RISK_RED
