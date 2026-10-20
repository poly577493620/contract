"""三种报告导出。"""

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from smart_compliance.domain.models import ReviewReport
from smart_compliance.reporting import export_docx, export_json, export_pdf


def test_export_reports(sample_report: ReviewReport, output_dir: Path) -> None:
    """JSON、DOCX、PDF 均应可重新打开且包含核心标题。"""
    json_path = export_json(sample_report, output_dir / "review.json")
    docx_path = export_docx(sample_report, output_dir / "review.docx")
    pdf_path = export_pdf(sample_report, output_dir / "review.pdf")
    assert ReviewReport.model_validate_json(json_path.read_text(encoding="utf-8")).sample_id == "005"
    assert "教学模拟/测试数据" in json_path.read_text(encoding="utf-8")
    document = Document(docx_path)
    assert any("采购合同智能质检报告" in paragraph.text for paragraph in document.paragraphs)
    assert any("运行与证据审计" in paragraph.text for paragraph in document.paragraphs)
    assert any(
        "教学模拟/测试数据" in paragraph.text
        for section in document.sections
        for paragraph in section.footer.paragraphs
    )
    audit_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "采购结果 PDF SHA-256" in audit_text
    assert "a" * 64 in audit_text
    assert "采购结果 OCR 文本 SHA-256" in audit_text
    reader = PdfReader(pdf_path)
    assert len(reader.pages) >= 1


def test_exports_remove_local_absolute_path(sample_report: ReviewReport, output_dir: Path) -> None:
    """所有公开格式都不得暴露本机盘符、用户名或工作区绝对路径。"""
    private_path = r"F:\Claude code\test4\outputs\evaluations\run-1\raw\006.txt"
    model = sample_report.model.model_copy(update={"raw_output_file": private_path})
    report = sample_report.model_copy(update={"model": model})

    json_path = export_json(report, output_dir / "safe.json")
    docx_path = export_docx(report, output_dir / "safe.docx")
    pdf_path = export_pdf(report, output_dir / "safe.pdf")

    json_text = json_path.read_text(encoding="utf-8")
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
    for exported_text in (json_text, docx_text, pdf_text):
        assert "F:\\Claude code" not in exported_text
        assert "outputs/evaluations/run-1/raw/006.txt" in exported_text


def test_report_summary_count_validation(sample_report: ReviewReport) -> None:
    """汇总与明细不一致时必须拒绝。"""
    payload = sample_report.model_dump()
    payload["summary"]["different_count"] = 0
    try:
        ReviewReport.model_validate(payload)
    except ValueError as error:
        assert "汇总计数与明细不一致" in str(error)
    else:
        raise AssertionError("无效汇总未被拒绝")


def test_issue_only_report_uses_coverage_label(sample_report: ReviewReport, output_dir: Path) -> None:
    """外部问题清单报告不得把未输出的一致项显示为零。"""
    payload = sample_report.model_dump()
    payload["summary"]["covered_clause_count"] = 14
    report = ReviewReport.model_validate(payload)
    document = Document(export_docx(report, output_dir / "coverage.docx"))
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "覆盖编码" in table_text
    assert "14" in table_text
